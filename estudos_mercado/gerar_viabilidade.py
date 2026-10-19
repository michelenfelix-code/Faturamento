"""
Gera a planilha de viabilidade das 3 oportunidades finalistas do levantamento de mercado:
  1. BPO financeiro e controladoria para PMEs
  2. Consultoria em riscos psicossociais (NR-1)
  3. Agência de cuidadores de idosos

Como usar:
  python gerar_viabilidade.py
  -> cria "viabilidade_oportunidades.xlsx" na mesma pasta deste script.

Para testar outras premissas, altere o dicionário OPORTUNIDADES abaixo e rode de novo,
ou edite direto as células amarelas da aba "Premissas" no Excel (tudo recalcula sozinho).
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SAIDA = Path(__file__).with_name("viabilidade_oportunidades.xlsx")
MESES = 36

# ---------------------------------------------------------------------------
# Premissas (todas são ESTIMATIVAS para validar com cotações locais)
# ---------------------------------------------------------------------------
# Cada premissa: (rótulo, unidade, formato, {aba: valor}, observação)
OPORTUNIDADES = ["BPO financeiro", "Consultoria NR-1", "Agência de cuidadores"]
ABAS_FLUXO = ["Fluxo BPO", "Fluxo NR-1", "Fluxo Cuidadores"]

PREMISSAS = [
    ("Investimento inicial", "R$", "moeda", [25000, 40000, 50000],
     "Estimativa. BPO: software, BI, marketing, CNPJ. NR-1: equipe técnica, software, material. "
     "Cuidadores: escritório, seguro, seleção/treinamento, marketing."),
    ("Ticket mensal recorrente por cliente", "R$/mês", "moeda", [1500, 1200, 5500],
     "Estimativa a validar com cotação oculta. Cuidadores: valor cobrado da família por um "
     "cuidador mensal (ex.: plantão 12x36 diurno)."),
    ("Receita de implantação por novo cliente", "R$", "moeda", [1000, 8000, 0],
     "Cobrança única na entrada. BPO: onboarding/organização inicial. "
     "NR-1: diagnóstico psicossocial + revisão do PGR."),
    ("Clientes no 1º mês", "clientes", "num", [2, 1, 2],
     "Estimativa conservadora (rede de contatos)."),
    ("Novos clientes por mês", "clientes/mês", "num", [2, 1.5, 1.5],
     "Estimativa. Ajustada pelo fator de cenário."),
    ("Capacidade máxima de clientes", "clientes", "num", [40, 30, 40],
     "Limite operacional da equipe prevista. Acima disso, contratar (rever custos fixos)."),
    ("Cancelamento mensal (churn)", "% ao mês", "pct", [0.02, 0.03, 0.05],
     "Estimativa. Cuidadores tem churn maior (internação, falecimento, mudança)."),
    ("Custo variável", "% da receita", "pct", [0.30, 0.35, 0.78],
     "BPO: horas de analista + software por cliente. NR-1: psicólogo parceiro e horas técnicas. "
     "Cuidadores: salário + encargos + benefícios do cuidador."),
    ("Impostos sobre a receita", "% da receita", "pct", [0.06, 0.06, 0.11],
     "Referência: Simples Nacional, faixa inicial. Confirme CNAE, anexo e Fator R. "
     "Cuidadores: cessão de mão de obra tem restrições no Simples, confirmar enquadramento."),
    ("Custos fixos mensais", "R$/mês", "moeda", [3000, 4000, 6000],
     "Estimativa: sistemas, contador, internet, marketing, aluguel/coworking."),
    ("Pró-labore mensal", "R$/mês", "moeda", [5000, 6000, 5000],
     "Retirada do(a) sócio(a). Ajuste para sua necessidade."),
]
# Linha na aba Premissas de cada premissa (começa na linha 5)
LINHA = {rotulo: 5 + i for i, (rotulo, *_ ) in enumerate(PREMISSAS)}

CENARIOS = [("Pessimista", 0.7, 0.9), ("Base", 1.0, 1.0), ("Otimista", 1.3, 1.1)]
TMA_MENSAL = 0.012

# ---------------------------------------------------------------------------
# Estilos
# ---------------------------------------------------------------------------
FONTE = "Arial"
AZUL = Font(name=FONTE, color="0000FF")
PRETO = Font(name=FONTE, color="000000")
VERDE = Font(name=FONTE, color="008000")
NEGRITO = Font(name=FONTE, bold=True)
TITULO = Font(name=FONTE, bold=True, size=14)
CAB = Font(name=FONTE, bold=True, color="FFFFFF")
FUNDO_CAB = PatternFill("solid", fgColor="1F4E78")
AMARELO = PatternFill("solid", fgColor="FFFF00")
CINZA = PatternFill("solid", fgColor="F2F2F2")
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))

FMT = {
    "moeda": 'R$ #,##0;(R$ #,##0);"-"',
    "num": '#,##0.0;(#,##0.0);"-"',
    "pct": '0.0%;(0.0%);"-"',
    "int": '0',
}


def cabecalho(ws, linha, textos):
    for c, t in enumerate(textos, start=1):
        cel = ws.cell(row=linha, column=c, value=t)
        cel.font = CAB
        cel.fill = FUNDO_CAB
        cel.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def larguras(ws, valores):
    for i, w in enumerate(valores, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


# ---------------------------------------------------------------------------
# Aba Leia-me
# ---------------------------------------------------------------------------
def aba_leiame(wb):
    ws = wb.active
    ws.title = "Leia-me"
    linhas = [
        ("Viabilidade das 3 oportunidades finalistas", TITULO),
        ("Levantamento de mercado: BPO financeiro, Consultoria NR-1 e Agência de cuidadores", NEGRITO),
        ("", None),
        ("COMO USAR", NEGRITO),
        ("1. Edite só as células AMARELAS com texto AZUL na aba 'Premissas'. Todo o resto é fórmula.", PRETO),
        ("2. Escolha o cenário (Pessimista / Base / Otimista) na célula B18 da aba 'Premissas'.", PRETO),
        ("3. Veja os indicadores comparados na aba 'Resumo' (payback, VPL, TIR, ponto de equilíbrio).", PRETO),
        ("4. As abas 'Fluxo ...' mostram mês a mês (36 meses) clientes, receita, custos e caixa acumulado.", PRETO),
        ("5. Use a aba 'Validação de campo' para registrar concorrentes e cotações no Alto Tietê "
         "e depois ajustar as premissas.", PRETO),
        ("", None),
        ("LEGENDA DE CORES", NEGRITO),
        ("Texto azul + fundo amarelo = premissa digitada (pode alterar)", AZUL),
        ("Texto preto = fórmula calculada (não alterar)", PRETO),
        ("Texto verde = vínculo com outra aba", VERDE),
        ("", None),
        ("ATENÇÃO", NEGRITO),
        ("Todas as premissas financeiras são ESTIMATIVAS iniciais, não dados de mercado. "
         "Valide preços, custos e impostos antes de qualquer decisão.", PRETO),
        ("Os impostos usam alíquota de referência do Simples Nacional. Confirme CNAE, anexo e Fator R.", PRETO),
        ("O modelo trata clientes como média (pode haver frações) e não inclui reajuste de preços "
         "nem capital de giro: é uma visão gerencial simplificada.", PRETO),
        ("", None),
        ("FONTES DOS DADOS DE DEMANDA (levantamento de out/2026)", NEGRITO),
        ("BPO / reforma tributária: https://www.imarcgroup.com/report/pt-br/brazil-business-process-outsourcing-market", PRETO),
        ("Contratações pela reforma (Robert Half): https://fenacon.org.br/reforma-tributaria/"
         "reforma-tributaria-impulsiona-contratacao-de-contadores-e-tributaristas/", PRETO),
        ("NR-1 em vigor desde 26/05/2026: https://www.rsdata.com.br/?p=36737", PRETO),
        ("Fiscalização NR-1: https://www.contabeis.com.br/noticias/76472/"
         "nr-1-entra-em-fase-de-fiscalizacao-e-pressiona-empresas/", PRETO),
        ("Afastamentos por transtornos mentais 2025: https://outraspalavras.net/outrasaude/"
         "quais-as-doencas-que-mais-afastaram-trabalhadores-em-2025/", PRETO),
        ("Cuidadores +74% (Sebrae): https://fenacon.org.br/noticias/"
         "em-cinco-anos-abertura-de-pequenos-negocios-de-servico-de-cuidadores-cresceu-74-no-brasil/", PRETO),
        ("Envelhecimento (Censo 2022): https://www.em.com.br/gerais/2026/07/"
         "7467827-com-o-pais-envelhecendo-quem-cuidara-dos-nossos-idosos-e-de-nos-mesmos.html", PRETO),
    ]
    for i, (texto, fonte) in enumerate(linhas, start=1):
        cel = ws.cell(row=i, column=1, value=texto)
        if fonte:
            cel.font = fonte
        if fonte is AZUL:
            cel.fill = AMARELO
    ws.column_dimensions["A"].width = 130


# ---------------------------------------------------------------------------
# Aba Premissas
# ---------------------------------------------------------------------------
def aba_premissas(wb):
    ws = wb.create_sheet("Premissas")
    ws["A1"] = "Premissas do modelo"
    ws["A1"].font = TITULO
    ws["A2"] = "Células amarelas = estimativas editáveis. Valores iniciais são hipóteses a validar."
    ws["A2"].font = Font(name=FONTE, italic=True)
    cabecalho(ws, 4, ["Premissa"] + OPORTUNIDADES + ["Unidade", "Observação"])

    for rotulo, unidade, fmt, valores, obs in PREMISSAS:
        r = LINHA[rotulo]
        ws.cell(row=r, column=1, value=rotulo).font = PRETO
        for c, v in enumerate(valores, start=2):
            cel = ws.cell(row=r, column=c, value=v)
            cel.font = AZUL
            cel.fill = AMARELO
            cel.number_format = FMT[fmt]
        ws.cell(row=r, column=5, value=unidade).font = PRETO
        o = ws.cell(row=r, column=6, value=obs)
        o.font = PRETO
        o.alignment = Alignment(wrap_text=True, vertical="top")

    # Cenários
    ws["A17"] = "Cenário"
    ws["A17"].font = NEGRITO
    ws["A18"] = "Cenário escolhido"
    ws["B18"] = "Base"
    ws["B18"].font = AZUL
    ws["B18"].fill = AMARELO
    dv = DataValidation(type="list", formula1='"Pessimista,Base,Otimista"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add("B18")
    ws["F18"] = "Lista suspensa: Pessimista, Base ou Otimista."

    cabecalho(ws, 20, ["Cenário", "Fator novos clientes", "Fator ticket"])
    for i, (nome, f_novos, f_ticket) in enumerate(CENARIOS, start=21):
        ws.cell(row=i, column=1, value=nome).font = PRETO
        for c, v in ((2, f_novos), (3, f_ticket)):
            cel = ws.cell(row=i, column=c, value=v)
            cel.font = AZUL
            cel.fill = AMARELO
            cel.number_format = "0.00x"
    ws["F21"] = "Fatores multiplicam 'Novos clientes por mês' e os tickets (recorrente e implantação)."

    ws["A25"] = "Fator novos clientes (ativo)"
    ws["B25"] = "=INDEX($B$21:$B$23,MATCH($B$18,$A$21:$A$23,0))"
    ws["A26"] = "Fator ticket (ativo)"
    ws["B26"] = "=INDEX($C$21:$C$23,MATCH($B$18,$A$21:$A$23,0))"
    for ref in ("B25", "B26"):
        ws[ref].number_format = "0.00x"
        ws[ref].font = PRETO

    ws["A28"] = "Taxa mínima de atratividade (TMA)"
    ws["B28"] = TMA_MENSAL
    ws["B28"].font = AZUL
    ws["B28"].fill = AMARELO
    ws["B28"].number_format = "0.00%"
    ws["E28"] = "% ao mês"
    ws["F28"] = ("Custo de oportunidade do dinheiro (ex.: rendimento de aplicação conservadora). "
                 "Usada no VPL. Ajuste para sua realidade.")

    for r in range(4, 29):
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = BORDA
    larguras(ws, [38, 18, 18, 22, 14, 90])
    ws.freeze_panes = "B5"
    return ws


# ---------------------------------------------------------------------------
# Abas de fluxo mensal
# ---------------------------------------------------------------------------
COLUNAS_FLUXO = [
    "Mês", "Clientes ativos", "Novos clientes (efetivos)", "Receita recorrente",
    "Receita de implantação", "Receita total", "Impostos", "Custos variáveis",
    "Margem de contribuição", "Custos fixos + pró-labore", "Resultado operacional",
    "Fluxo de caixa", "Caixa acumulado", "Payback atingido? (1=sim)",
]
L0 = 5                # linha do mês 0
LFIM = L0 + MESES     # linha do mês 36


def p(col, rotulo):
    """Referência absoluta a uma premissa na aba Premissas."""
    return f"Premissas!${col}${LINHA[rotulo]}"


def aba_fluxo(wb, nome_aba, nome_op, col):
    ws = wb.create_sheet(nome_aba)
    ws["A1"] = f"Fluxo mensal - {nome_op}"
    ws["A1"].font = TITULO
    ws["A2"] = "Cenário ativo:"
    ws["B2"] = "=Premissas!$B$18"
    ws["B2"].font = VERDE
    cabecalho(ws, 4, COLUNAS_FLUXO)

    cap = p(col, "Capacidade máxima de clientes")
    churn = p(col, "Cancelamento mensal (churn)")
    inicial = p(col, "Clientes no 1º mês")
    novos = p(col, "Novos clientes por mês")
    ticket = p(col, "Ticket mensal recorrente por cliente")
    impl = p(col, "Receita de implantação por novo cliente")
    imp = p(col, "Impostos sobre a receita")
    cv = p(col, "Custo variável")
    fixos = p(col, "Custos fixos mensais")
    pro = p(col, "Pró-labore mensal")
    inv = p(col, "Investimento inicial")
    f_novos, f_ticket = "Premissas!$B$25", "Premissas!$B$26"

    # Mês 0: só o investimento
    ws.cell(row=L0, column=1, value=0)
    for c in range(2, 12):
        ws.cell(row=L0, column=c, value=0)
    ws.cell(row=L0, column=12, value=f"=-{inv}")
    ws.cell(row=L0, column=13, value=f"=L{L0}")
    ws.cell(row=L0, column=14, value=f"=IF(M{L0}>=0,1,0)")

    for r in range(L0 + 1, LFIM + 1):
        a = r - 1  # linha anterior
        f = {
            1: f"=A{a}+1",
            2: f"=MIN({cap},B{a}*(1-{churn})+IF(A{r}=1,{inicial},{novos}*{f_novos}))",
            3: f"=MAX(0,B{r}-B{a}*(1-{churn}))",
            4: f"=B{r}*{ticket}*{f_ticket}",
            5: f"=C{r}*{impl}*{f_ticket}",
            6: f"=D{r}+E{r}",
            7: f"=F{r}*{imp}",
            8: f"=F{r}*{cv}",
            9: f"=F{r}-G{r}-H{r}",
            10: f"={fixos}+{pro}",
            11: f"=I{r}-J{r}",
            12: f"=K{r}",
            13: f"=M{a}+L{r}",
            14: f"=IF(M{r}>=0,1,0)",
        }
        for c, formula in f.items():
            ws.cell(row=r, column=c, value=formula)

    for r in range(L0, LFIM + 1):
        for c in range(1, 15):
            cel = ws.cell(row=r, column=c)
            if c in (2, 3):
                cel.number_format = FMT["num"]
            elif c in (1, 14):
                cel.number_format = FMT["int"]
            else:
                cel.number_format = FMT["moeda"]
            if cel.font != VERDE:
                cel.font = PRETO
            if (r - L0) % 12 == 0 and r != L0:
                cel.fill = CINZA
    ws.cell(row=L0, column=12).font = VERDE

    ws.cell(row=LFIM + 2, column=1,
            value="Linhas cinza = fim de cada ano (meses 12, 24 e 36).").font = Font(name=FONTE, italic=True)
    larguras(ws, [7, 12, 14, 16, 16, 16, 14, 16, 16, 16, 16, 16, 16, 12])
    ws.row_dimensions[4].height = 45
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------------------
# Aba Resumo
# ---------------------------------------------------------------------------
def aba_resumo(wb):
    ws = wb.create_sheet("Resumo", 1)
    ws["A1"] = "Resumo comparativo - 36 meses"
    ws["A1"].font = TITULO
    ws["A2"] = "Cenário ativo:"
    ws["B2"] = "=Premissas!$B$18"
    ws["B2"].font = VERDE
    cabecalho(ws, 4, ["Indicador"] + OPORTUNIDADES + ["Como ler"])

    l12, l24, l36 = L0 + 12, L0 + 24, LFIM
    indicadores = [
        ("Investimento inicial", "moeda",
         lambda a, c: f"={p(c, 'Investimento inicial')}", "Valor aplicado no mês 0."),
        ("Clientes ativos - mês 12", "num", lambda a, c: f"='{a}'!B{l12}", ""),
        ("Clientes ativos - mês 36", "num", lambda a, c: f"='{a}'!B{l36}", ""),
        ("Receita mensal - mês 12", "moeda", lambda a, c: f"='{a}'!F{l12}", ""),
        ("Receita mensal - mês 36", "moeda", lambda a, c: f"='{a}'!F{l36}", ""),
        ("Resultado mensal - mês 12", "moeda", lambda a, c: f"='{a}'!K{l12}",
         "Depois de impostos, custos e pró-labore."),
        ("Resultado mensal - mês 36", "moeda", lambda a, c: f"='{a}'!K{l36}", ""),
        ("Margem líquida - mês 12", "pct",
         lambda a, c: f"=IFERROR('{a}'!K{l12}/'{a}'!F{l12},0)", "Resultado / receita."),
        ("Margem líquida - mês 36", "pct",
         lambda a, c: f"=IFERROR('{a}'!K{l36}/'{a}'!F{l36},0)", ""),
        ("Receita acumulada - 36 meses", "moeda",
         lambda a, c: f"=SUM('{a}'!F{L0}:F{LFIM})", ""),
        ("Caixa acumulado - mês 36", "moeda", lambda a, c: f"='{a}'!M{l36}",
         "Já descontado o investimento inicial."),
        ("Ponto de equilíbrio (clientes)", "num",
         lambda a, c: (f"=IFERROR(({p(c, 'Custos fixos mensais')}+{p(c, 'Pró-labore mensal')})/"
                       f"({p(c, 'Ticket mensal recorrente por cliente')}*Premissas!$B$26*"
                       f"(1-{p(c, 'Impostos sobre a receita')}-{p(c, 'Custo variável')})),\"n/d\")"),
         "Clientes recorrentes para cobrir custos fixos e pró-labore (sem implantação)."),
        ("Payback (meses)", "int",
         lambda a, c: f"=IFERROR(MATCH(1,'{a}'!N{L0}:N{LFIM},0)-1,\"Acima de 36\")",
         "Mês em que o caixa acumulado volta a ficar positivo."),
        ("VPL - 36 meses", "moeda",
         lambda a, c: f"=NPV(Premissas!$B$28,'{a}'!L{L0 + 1}:L{LFIM})+'{a}'!L{L0}",
         "Positivo = rende mais que a TMA."),
        ("TIR mensal", "pct", lambda a, c: f"=IFERROR(IRR('{a}'!L{L0}:L{LFIM}),\"n/d\")",
         "Comparar com a TMA."),
        ("TIR anual equivalente", "pct",
         lambda a, c: f"=IFERROR((1+IRR('{a}'!L{L0}:L{LFIM}))^12-1,\"n/d\")", ""),
    ]
    colunas_premissa = ["B", "C", "D"]
    for i, (rotulo, fmt, formula, como_ler) in enumerate(indicadores, start=5):
        ws.cell(row=i, column=1, value=rotulo).font = PRETO
        for j, (aba, colp) in enumerate(zip(ABAS_FLUXO, colunas_premissa), start=2):
            cel = ws.cell(row=i, column=j, value=formula(aba, colp))
            cel.font = VERDE
            cel.number_format = FMT[fmt]
            cel.alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=5, value=como_ler).font = PRETO
        for c in range(1, 6):
            ws.cell(row=i, column=c).border = BORDA

    nota = 5 + len(indicadores) + 1
    ws.cell(row=nota, column=1,
            value="Todos os valores dependem das premissas estimadas; troque o cenário em Premissas!B18 "
                  "para testar sensibilidade.").font = Font(name=FONTE, italic=True)
    larguras(ws, [34, 18, 18, 22, 70])
    ws.freeze_panes = "B5"


# ---------------------------------------------------------------------------
# Aba Validação de campo
# ---------------------------------------------------------------------------
def aba_validacao(wb):
    ws = wb.create_sheet("Validação de campo")
    ws["A1"] = "Validação de campo - concorrentes e cotações no Alto Tietê"
    ws["A1"].font = TITULO
    ws["A2"] = ("Preencha uma linha por concorrente encontrado (Google Maps, indicações, cotação oculta). "
                "A linha 5 é um EXEMPLO fictício de formato: substitua pelos seus dados.")
    ws["A2"].font = Font(name=FONTE, italic=True)
    cols = ["Oportunidade", "Cidade", "Empresa", "Nota Google", "Nº avaliações",
            "Preço cotado (R$)", "Prazo para atender (dias)", "Pontos fracos / reclamações", "Data"]
    cabecalho(ws, 4, cols)
    exemplo = ["BPO financeiro", "Mogi das Cruzes", "Exemplo Ltda (fictício)", 4.3, 27, 1800, 10,
               "Só faz lançamentos, sem relatório gerencial", "2026-10-20"]
    for c, v in enumerate(exemplo, start=1):
        cel = ws.cell(row=5, column=c, value=v)
        cel.font = Font(name=FONTE, color="0000FF", italic=True)
        cel.fill = AMARELO
    ws["F5"].number_format = FMT["moeda"]

    dv = DataValidation(type="list", formula1=f'"{",".join(OPORTUNIDADES)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("A5:A200")

    linha_res = 4
    ws.cell(row=linha_res - 1, column=11, value="Resumo por oportunidade").font = NEGRITO
    cabecalho_res = ["Oportunidade", "Concorrentes", "Nota média", "Preço médio (R$)", "Prazo médio (dias)"]
    for c, t in enumerate(cabecalho_res, start=11):
        cel = ws.cell(row=linha_res, column=c, value=t)
        cel.font = CAB
        cel.fill = FUNDO_CAB
        cel.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, op in enumerate(OPORTUNIDADES, start=linha_res + 1):
        ws.cell(row=i, column=11, value=op).font = PRETO
        ws.cell(row=i, column=12, value=f'=COUNTIF($A$5:$A$200,K{i})')
        ws.cell(row=i, column=13, value=f'=IFERROR(AVERAGEIF($A$5:$A$200,K{i},$D$5:$D$200),"-")')
        ws.cell(row=i, column=14, value=f'=IFERROR(AVERAGEIF($A$5:$A$200,K{i},$F$5:$F$200),"-")')
        ws.cell(row=i, column=15, value=f'=IFERROR(AVERAGEIF($A$5:$A$200,K{i},$G$5:$G$200),"-")')
        ws.cell(row=i, column=13).number_format = "0.0"
        ws.cell(row=i, column=14).number_format = FMT["moeda"]
        ws.cell(row=i, column=15).number_format = "0.0"
        for c in range(12, 16):
            ws.cell(row=i, column=c).font = PRETO
    ws.cell(row=linha_res + 5, column=11,
            value="Use o preço médio para revisar o ticket na aba Premissas.").font = Font(name=FONTE, italic=True)
    ws["F4"].comment = Comment("Valor mensal cobrado pelo concorrente para o serviço equivalente.", "Modelo")

    larguras(ws, [22, 18, 26, 11, 12, 16, 14, 42, 12, 3, 22, 13, 11, 16, 14])
    ws.row_dimensions[4].height = 32
    ws.freeze_panes = "A5"


def main():
    wb = Workbook()
    aba_leiame(wb)
    aba_premissas(wb)
    for aba, op, col in zip(ABAS_FLUXO, OPORTUNIDADES, ["B", "C", "D"]):
        aba_fluxo(wb, aba, op, col)
    aba_resumo(wb)
    aba_validacao(wb)
    wb.save(SAIDA)
    print(f"Planilha gerada: {SAIDA}")


if __name__ == "__main__":
    main()
